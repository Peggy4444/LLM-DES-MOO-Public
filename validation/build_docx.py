"""Generate the short participant packet as a Word document.

Design intent (revised):
- Total writing time for the participant should be ~15 minutes.
- Everything else in the session is verbal discussion (recorded).
- Only include items that directly feed a table or a number in the paper.
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor

OUT = Path(__file__).parent / "Participant_Packet.docx"


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return h


def add_para(doc, text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p


def add_bullet(doc, text):
    return doc.add_paragraph(text, style="List Bullet")


def add_checkbox_line(doc, text):
    p = doc.add_paragraph()
    p.add_run("☐  ").font.size = Pt(12)
    p.add_run(text).font.size = Pt(11)
    return p


def add_blank_line(doc, char="_", n=70):
    add_para(doc, char * n)


def add_field(doc, label, width_chars=60):
    p = doc.add_paragraph()
    r = p.add_run(f"{label} ")
    r.font.size = Pt(11)
    r.bold = True
    p.add_run("_" * width_chars).font.size = Pt(11)


def add_table_header(doc, headers, col_widths_cm=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Light Grid Accent 1"
    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
    if col_widths_cm:
        for row in tbl.rows:
            for i, w in enumerate(col_widths_cm):
                row.cells[i].width = Cm(w)
    return tbl


def add_row(tbl, values, col_widths_cm=None):
    row = tbl.add_row().cells
    for i, v in enumerate(values):
        row[i].text = ""
        run = row[i].paragraphs[0].add_run(str(v))
        run.font.size = Pt(10)
    if col_widths_cm:
        for i, w in enumerate(col_widths_cm):
            row[i].width = Cm(w)
    return row


def build():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ============ Cover / consent (page 1) ============
    title = doc.add_heading(
        "Participant Packet — External Evaluation", level=0
    )
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    add_para(
        doc,
        "An LLM-based decision-support framework for simulation-driven "
        "production-line optimization.",
        italic=True,
    )
    add_para(doc, "Facilitator: Pegah Rahimian, Uppsala University "
                  "(rahimian.pegah@gmail.com)")
    add_para(doc, "")
    add_field(doc, "Date:", width_chars=25)
    add_field(doc, "Participant ID:", width_chars=10)

    add_para(doc, "")
    add_heading(doc, "Consent", level=1)
    add_para(
        doc,
        "You are invited to a ~90-minute session in which you will try a "
        "research prototype and share your professional opinion on its "
        "outputs. Screen and audio will be recorded with your permission and "
        "deleted after publication. Your comments may be quoted anonymously "
        "in the paper by role and industry sector only. You may stop or "
        "withdraw at any time. No production data from your company is used.",
    )
    add_checkbox_line(doc, "I consent to participate.")
    add_checkbox_line(doc, "I consent to screen and audio recording.")
    add_checkbox_line(doc, "I consent to being quoted anonymously "
                           "by role and sector.")
    add_field(doc, "Signature:", width_chars=45)

    doc.add_page_break()

    # ============ Background (page 2, top) ============
    add_heading(doc, "About you (short)", level=1)
    add_para(doc, "Please answer briefly. Skip anything you prefer not to answer.",
             italic=True)
    add_para(doc, "")

    add_field(doc, "1. Role:", width_chars=55)
    add_field(doc, "2. Sector (truck / cars / components / other):",
              width_chars=30)
    add_field(doc, "3. Years in simulation or production engineering:",
              width_chars=8)
    add_para(doc, "4. Simulation tools used in the past 3 years (circle):")
    add_para(doc, "    FACTS  ·  Plant Simulation  ·  Arena  ·  AnyLogic  ·  "
                  "SimPy  ·  Other: ______________________________")
    add_para(doc,
             "5. Could your team send production data to a cloud LLM "
             "(OpenAI, Anthropic, Google) today?")
    add_para(doc, "    Yes  ·  No  ·  Unsure  ·  Only with anonymization")

    add_para(doc, "")

    # ============ Per-artifact rating (page 2, main) ============
    add_heading(doc, "Rating of the system outputs", level=1)
    add_para(
        doc,
        "For each artifact the framework produced, rate three dimensions on "
        "a 1–5 scale (1 = very poor, 5 = excellent). Leave the comment blank "
        "if you have nothing to add.",
    )
    add_para(doc, "Correctness — matches the underlying data or system.  "
                  "Trust — you would rely on it for a decision.  "
                  "Actionability — you or your team could act on it.",
             italic=True)
    add_para(doc, "")

    artifacts = [
        "Topology diagram (visualizer)",
        "Pareto front and knee selection (optimizer)",
        "Bottleneck ranking (SCORE)",
        "Bottleneck natural-language rationale",
        "Model-adaptation KPI report",
    ]
    widths = [6.5, 1.8, 1.5, 2.0, 5.2]
    tbl = add_table_header(
        doc,
        ["Artifact", "Correct", "Trust", "Actionable", "Optional comment"],
        col_widths_cm=widths,
    )
    for a in artifacts:
        add_row(tbl, [a, "", "", "", ""], col_widths_cm=widths)

    add_para(doc, "")

    # Overall three-item system rating - the three numbers most useful in text.
    add_heading(doc, "Overall system (three quick items)", level=2)
    add_para(doc, "Rate 1 (strongly disagree) to 5 (strongly agree):",
             italic=True)
    overall = [
        "S1. I would use this tool in my work.",
        "S2. I trust the framework's outputs enough to act on them.",
        "S3. This tool could reduce the expertise barrier for a "
        "non-specialist colleague.",
    ]
    widths3 = [12.5, 1.0, 1.0, 1.0, 1.0, 1.0]
    tbl3 = add_table_header(
        doc, ["Statement", "1", "2", "3", "4", "5"], col_widths_cm=widths3
    )
    for s in overall:
        add_row(tbl3, [s, "", "", "", "", ""], col_widths_cm=widths3)

    doc.add_page_break()

    # ============ Paired interpretation (page 3) ============
    add_heading(doc, "Paired interpretation tasks", level=1)
    add_para(
        doc,
        "For each of the three tasks, the facilitator will first show raw "
        "data and ask you to interpret it aloud. The LLM's interpretation is "
        "revealed afterwards. Only a short written mark is needed here; the "
        "rest is discussion.",
    )
    add_para(doc, "")

    widths_p = [3.5, 8.5, 5.5]
    tblp = add_table_header(
        doc,
        [
            "Task",
            "Did the LLM interpretation agree with yours?",
            "One-line note (what it got right / missed / got wrong)",
        ],
        col_widths_cm=widths_p,
    )
    add_row(
        tblp,
        [
            "T1 — Topology",
            "Full  ·  Partial  ·  None",
            "",
        ],
        col_widths_cm=widths_p,
    )
    add_row(
        tblp,
        [
            "T2 — Pareto front",
            "Full  ·  Partial  ·  None",
            "",
        ],
        col_widths_cm=widths_p,
    )
    add_row(
        tblp,
        [
            "T3 — Bottleneck",
            "Full  ·  Partial  ·  None",
            "",
        ],
        col_widths_cm=widths_p,
    )

    add_para(doc, "")
    add_para(doc,
             "Space for your notes during the discussion (optional):",
             italic=True)
    for _ in range(8):
        add_blank_line(doc)

    doc.add_page_break()

    # ============ Open reflection (page 4) - the important page ============
    add_heading(doc, "Your reflections", level=1)
    add_para(
        doc,
        "This is the most important part for us. Please answer freely — "
        "we will discuss each of these verbally.",
        italic=True,
    )
    add_para(doc, "")

    prompts = [
        "1. What is the single most useful output the framework produced today?",
        "2. What would you not trust yet, and what would have to change?",
        "3. Under your company's IT and data-governance rules, what would need "
        "to be true before you could use this on a real production line?",
        "4. Anything else you would like the researchers to know?",
    ]
    for prompt in prompts:
        add_para(doc, prompt, bold=True)
        for _ in range(4):
            add_blank_line(doc)
        add_para(doc, "")

    add_field(doc, "Preferred email for the 1-page summary we will send you:",
              width_chars=40)

    add_para(doc, "")
    add_para(
        doc,
        "Thank you. We will send you a short written summary of what we heard "
        "within 48 hours and ask you to confirm it reflects your view "
        "correctly before we use it in the paper.",
        italic=True,
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
